#!/usr/bin/env python3
"""Builds the teacher's evaluation workbook template (upload to Google Drive → opens as a Google Sheet).

Usage:
  npx tsx scripts/station-catalog.ts > catalog.json
  python3 scripts/build-evaluation-sheet.py catalog.json out.xlsx [sample-tabs.json]

Raw tabs (Attempts, Station Responses, …) are filled by the app's "Update Google Sheet" button
(Edge Function sheet-export). Analysis tabs use only classic functions (COUNTIFS / AVERAGEIFS / IFERROR)
so they survive the .xlsx → Google Sheets conversion. Teacher tabs (Manual Exclusions, Handover Scoring)
are never overwritten by the export.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

catalog = json.load(open(sys.argv[1], encoding="utf-8"))
out_path = sys.argv[2]
sample = json.load(open(sys.argv[3], encoding="utf-8")) if len(sys.argv) > 3 else {}
TABS, TEACHER = catalog["tabs"], catalog["teacherTabs"]
N = int(os.environ.get("EVAL_ROWS", "5000"))  # analysed rows per raw tab

NAVY, TEAL, GREY = "12324A", "2B7279", "F2F4F7"
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor=NAVY)
input_fill = PatternFill("solid", fgColor="FFF8E1")
title_font = Font(bold=True, size=16, color=NAVY)
note_font = Font(italic=True, color="555555")
red, amber, green = PatternFill("solid", fgColor="F8D7DA"), PatternFill("solid", fgColor="FFF3CD"), PatternFill("solid", fgColor="D4EDDA")

wb = Workbook()
wb.remove(wb.active)


def sheet(name, headers=None, widths=None, teacher=False):
    ws = wb.create_sheet(name)
    if headers:
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = PatternFill("solid", fgColor="8A5A00") if teacher else header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        ws.row_dimensions[1].height = 42
    for index, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    return ws


def col(tab, header):
    """Column letter of a header in a raw tab (shared with the app's export schema)."""
    return get_column_letter(TABS[tab]["headers"].index(header) + 1)


def rng(tab, header):
    return f"'{TABS[tab]['name']}'!${col(tab, header)}$2:${col(tab, header)}${N}"


A = lambda h: rng("attempts", h)
R = lambda h: rng("responses", h)
F = lambda h: rng("feedback", h)
RO = lambda h: rng("roster", h)

# ---------------------------------------------------------------- README
readme = sheet("README", widths=[30, 110])
readme["A1"] = "Pelvic Fx Game — Evaluation workbook"
readme["A1"].font = title_font
rows = [
    ("Content version", catalog["contentVersion"]),
    ("How it fills", "Teacher → app → Faculty → 'Update Google Sheet now'. Raw tabs are replaced each time; analysis tabs recalculate. Manual Exclusions and Handover Scoring are yours and are never overwritten."),
    ("Who is counted", "Statistics use only rows where 'Exclude from analysis' is FALSE. QA/DEMO/TEST IDs, practice attempts and anyone in Manual Exclusions are excluded automatically but stay visible."),
    ("Learning objectives", "LO1 mechanism & associated injuries · LO2 history & examination · LO3 imaging choice & interpretation · LO4 open vs closed fracture · LO5 initial treatment (shock, resuscitation, open fracture) · LO6 pelvic binder · LO7 referral & handover."),
    ("Design", "Pre-test (8 items) → 24 practice stations with feedback and correction → boss reviews → post-test (8 items, parallel form). Formative, supervised learning; not a certification of competence."),
    ("Pass standard", "Post-test ≥ 6/8 correct on the first try AND both must-pass safety items correct (FS2 no Foley with meatal blood; FS6 binder at the greater trochanters). A conjunctive standard: high marks cannot compensate for a safety error."),
    ("Raw gain", "Post-test score − pre-test score (same 8-item blueprint, other form)."),
    ("Normalized gain (g)", "Hake: (post − pre) / (8 − pre). 0.3–0.7 = medium, > 0.7 = high. Fairer than raw gain when pre-test scores differ."),
    ("Paired t / p / dz", "Pre-Post by Cohort: paired t-test of gain (df = n − 1), two-sided p from TDIST, Cohen's dz = mean gain / SD of gain (0.2 small, 0.5 medium, 0.8 large)."),
    ("Difficulty p", "Item Analysis: proportion correct on the first try. < 0.30 = very hard (check the teaching or the item); > 0.90 = very easy."),
    ("Discrimination D", "First-try rate of the upper 27% minus the lower 27% of learners by post-test score. D < 0.20 suggests the item does not separate stronger and weaker learners. Needs at least 4 included learners."),
    ("Bloom level", "Station task type: remember · understand · apply · analyse. Use it to check items do not all test recall."),
    ("Kirkpatrick", "Level 1 reaction = Course Feedback (1–5). Level 2 learning = pre/post and objective mastery. Level 3 behaviour needs ward observation and is outside this game."),
    ("Misconceptions", "Specific wrong cards chosen on the first try, ranked by % of learners — use the top rows to plan the in-class discussion."),
    ("Handover scoring", "Score each learner's written handover reasoning with the SBAR rubric (0 = missing, 1 = partly, 2 = clear) in Handover Scoring; Total /8 calculates."),
    ("Privacy", "Contains student identifiers and free text. Keep the sheet private (never 'anyone with the link'); share only with course teachers. Delete or anonymise at the end of the retention period."),
]
for index, (label, text) in enumerate(rows, start=3):
    readme.cell(index, 1, label).font = Font(bold=True, color=NAVY)
    readme.cell(index, 2, text).alignment = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------- Dashboard
dash = sheet("Dashboard", widths=[44, 18, 18, 18, 18])
dash["A1"] = "Class dashboard"
dash["A1"].font = title_font
dash["A2"] = "Included learners only (Exclude = FALSE). Recalculates automatically."
dash["A2"].font = note_font
kpis = [
    ("Enrolled", f"=COUNTIFS({RO('Exclude from analysis')},FALSE)"),
    ("Started", f"=COUNTIFS({RO('Started')},\"yes\",{RO('Exclude from analysis')},FALSE)"),
    ("Completed the shift", f"=COUNTIFS({RO('Completed')},\"yes\",{RO('Exclude from analysis')},FALSE)"),
    ("Finished the post-test", f"=COUNTIFS({A('Post-test /8')},\">=0\",{A('Exclude from analysis')},FALSE)"),
    ("Met the pass standard", f"=COUNTIFS({A('Pass standard met')},TRUE,{A('Exclude from analysis')},FALSE)"),
    ("Pass rate", "=IFERROR(B8/B7,\"\")"),
    ("Mean pre-test /8", f"=IFERROR(AVERAGEIFS({A('Pre-test /8')},{A('Pre-test /8')},\">=0\",{A('Exclude from analysis')},FALSE),\"\")"),
    ("Mean post-test /8", f"=IFERROR(AVERAGEIFS({A('Post-test /8')},{A('Post-test /8')},\">=0\",{A('Exclude from analysis')},FALSE),\"\")"),
    ("Mean normalized gain (g)", f"=IFERROR(AVERAGEIFS({A('Normalized gain (g)')},{A('Normalized gain (g)')},\">=-1\",{A('Exclude from analysis')},FALSE),\"\")"),
    ("Mean time on task (min)", f"=IFERROR(AVERAGEIFS({A('Time on task (min)')},{A('Time on task (min)')},\">0\",{A('Exclude from analysis')},FALSE),\"\")"),
    ("Course feedback — useful (1–5)", f"=IFERROR(AVERAGEIFS({F('Useful for learning (1-5)')},{F('Exclude from analysis')},FALSE),\"\")"),
    ("Course feedback — enjoyable (1–5)", f"=IFERROR(AVERAGEIFS({F('Enjoyable (1-5)')},{F('Exclude from analysis')},FALSE),\"\")"),
    ("Course feedback — confidence (1–5)", f"=IFERROR(AVERAGEIFS({F('Confidence (1-5)')},{F('Exclude from analysis')},FALSE),\"\")"),
]
for index, (label, formula) in enumerate(kpis, start=4):
    dash.cell(index, 1, label).font = Font(bold=True)
    dash.cell(index, 2, formula)
dash["B9"].number_format = "0%"
for cell in ("B10", "B11", "B12", "B13", "B14", "B15", "B16"):
    dash[cell].number_format = "0.00"
dash["A18"] = "See: Objective Summary · Item Analysis · Pre-Post by Cohort · Misconceptions"
dash["A18"].font = note_font

# ---------------------------------------------------------------- Objective Summary
objectives = [
    ("LO1", "Mechanism of injury and associated injuries"),
    ("LO2", "History and physical examination"),
    ("LO3", "Choose and interpret imaging"),
    ("LO4", "Open vs closed pelvic fracture"),
    ("LO5", "Initial treatment: shock, resuscitation, open fracture"),
    ("LO6", "Pelvic binder: indication and application"),
    ("LO7", "Referral and handover"),
]
obj = sheet("Objective Summary", ["Objective", "Description", "Practice stations", "Practice responses (n)", "Practice first-try %", "Corrected after feedback %", "Pre-test first-try %", "Post-test first-try %", "Change (post − pre)"], [10, 46, 12, 14, 14, 16, 14, 14, 14])
for index, (lo, text) in enumerate(objectives, start=2):
    stations = sum(1 for s in catalog["stations"] if s["phase"] == "practice" and lo in s["objectives"].split(";"))
    like = f'"*{lo}*"'
    obj.append([
        lo, text, stations,
        f"=COUNTIFS({R('Objectives')},{like},{R('Phase')},\"practice\",{R('Exclude from analysis')},FALSE)",
        f"=IFERROR(AVERAGEIFS({R('First-try correct (1/0)')},{R('Objectives')},{like},{R('Phase')},\"practice\",{R('Exclude from analysis')},FALSE),\"\")",
        f"=IFERROR(COUNTIFS({R('Objectives')},{like},{R('Phase')},\"practice\",{R('First-try correct (1/0)')},0,{R('Corrected after feedback (1/0)')},1,{R('Exclude from analysis')},FALSE)/COUNTIFS({R('Objectives')},{like},{R('Phase')},\"practice\",{R('First-try correct (1/0)')},0,{R('Exclude from analysis')},FALSE),\"\")",
        f"=IFERROR(AVERAGEIFS({R('First-try correct (1/0)')},{R('Objectives')},{like},{R('Phase')},\"pre-test\",{R('Exclude from analysis')},FALSE),\"\")",
        f"=IFERROR(AVERAGEIFS({R('First-try correct (1/0)')},{R('Objectives')},{like},{R('Phase')},\"post-test\",{R('Exclude from analysis')},FALSE),\"\")",
        f"=IFERROR(H{index}-G{index},\"\")",
    ])
    for letter in "EFGHI":
        obj[f"{letter}{index}"].number_format = "0%"
for letter in "EGH":
    obj.conditional_formatting.add(f"{letter}2:{letter}8", ColorScaleRule(start_type="num", start_value=0.5, start_color="F8D7DA", mid_type="num", mid_value=0.65, mid_color="FFF3CD", end_type="num", end_value=0.8, end_color="D4EDDA"))
obj["A10"] = "Mastery colour: red < 50%, amber 50–79%, green ≥ 80% first try. Per-learner objective % is in Attempts (LO1–LO7 columns)."
obj["A10"].font = note_font

# ---------------------------------------------------------------- Item Analysis
items = sheet("Item Analysis", ["Station", "Phase", "Case", "Title", "Objectives", "Bloom level", "Must-pass", "Answered (n)", "Difficulty p (first-try)", "Corrected after feedback %", "Discrimination D", "Mean time (s)", "Hint use %", "Flag", "Teaching point"], [8, 11, 16, 34, 12, 11, 9, 10, 12, 13, 13, 10, 10, 30, 60])
for index, station in enumerate(catalog["stations"], start=2):
    sid = f'"{station["id"]}"'
    base = f"{R('Station')},{sid},{R('Exclude from analysis')},FALSE"
    items.append([
        station["id"], station["phase"], station["case"], station["title"], station["objectives"], station["bloom"], station["mustPass"],
        f"=COUNTIFS({base})",
        f"=IFERROR(AVERAGEIFS({R('First-try correct (1/0)')},{base}),\"\")",
        f"=IFERROR(COUNTIFS({base},{R('First-try correct (1/0)')},0,{R('Corrected after feedback (1/0)')},1)/COUNTIFS({base},{R('First-try correct (1/0)')},0),\"\")",
        f"=IFERROR(AVERAGEIFS({R('First-try correct (1/0)')},{base},{R('Post-test group')},\"upper\")-AVERAGEIFS({R('First-try correct (1/0)')},{base},{R('Post-test group')},\"lower\"),\"\")",
        f"=IFERROR(AVERAGEIFS({R('Time (s)')},{base},{R('Time (s)')},\">0\"),\"\")",
        f"=IFERROR(COUNTIFS({base},{R('Hints opened')},\">0\")/H{index},\"\")",
        f"=IF(H{index}<5,\"too few answers\",IF(I{index}<0.3,\"very hard — review teaching or item\",IF(I{index}>0.9,\"very easy\",IF(AND(K{index}<>\"\",K{index}<0.2),\"low discrimination — review item\",\"OK\"))))",
        station["key"],
    ])
    for letter in "IJM":
        items[f"{letter}{index}"].number_format = "0%"
    items[f"K{index}"].number_format = "0.00"
    items[f"L{index}"].number_format = "0"
last = len(catalog["stations"]) + 1
items.conditional_formatting.add(f"I2:I{last}", CellIsRule(operator="lessThan", formula=["0.3"], fill=red))
items.conditional_formatting.add(f"I2:I{last}", CellIsRule(operator="greaterThan", formula=["0.9"], fill=amber))
items.conditional_formatting.add(f"N2:N{last}", FormulaRule(formula=[f'AND(N2<>"OK",N2<>"too few answers")'], fill=amber))

# ---------------------------------------------------------------- raw tabs (filled by the export)
widths = {
    "cohorts": [24, 10, 10, 12, 10, 10, 10, 10, 10, 6, 12, 12, 14, 10],
    "misconceptions": [8, 11, 12, 46, 12, 12, 10, 70],
    "attempts": [22, 14, 14, 14, 18, 10, 16, 16, 16, 9, 9, 10, 10, 8, 10, 10, 10, 10, 10, 8] + [9] * 7 + [8, 10, 26],
    "responses": [22, 14, 14, 14, 11, 8, 18, 12, 11, 9, 10, 10, 9, 40, 8, 8, 10, 16, 10, 10],
    "handovers": [22, 14, 14, 14, 18, 60, 12, 16, 10],
    "reviews": [22, 14, 14, 14, 16, 16, 14, 14, 14, 40, 40, 40, 14],
    "feedback": [22, 14, 14, 14, 16, 10, 10, 10, 60, 10],
    "roster": [14, 14, 9, 10, 10, 26],
    "log": [24, 90],
}
order = ["cohorts", "misconceptions", "attempts", "responses", "handovers"]
for key in order:
    ws = sheet(TABS[key]["name"], TABS[key]["headers"], widths.get(key))
    for row in sample.get(TABS[key]["name"], [])[1:]:
        ws.append(row)
attempts_ws = wb[TABS["attempts"]["name"]]
lo_first, lo_last = col("attempts", "LO1 first-try %"), col("attempts", "LO7 first-try %")
attempts_ws.conditional_formatting.add(f"{lo_first}2:{lo_last}{N}", ColorScaleRule(start_type="num", start_value=50, start_color="F8D7DA", mid_type="num", mid_value=65, mid_color="FFF3CD", end_type="num", end_value=80, end_color="D4EDDA"))
pass_col = col("attempts", "Pass standard met")
attempts_ws.conditional_formatting.add(f"{pass_col}2:{pass_col}{N}", CellIsRule(operator="equal", formula=["FALSE"], fill=red))
attempts_ws.conditional_formatting.add(f"{pass_col}2:{pass_col}{N}", CellIsRule(operator="equal", formula=["TRUE"], fill=green))

# Teacher-owned: Handover Scoring (SBAR rubric 0–2 each).
scoring = sheet(TEACHER["handoverScoring"]["name"], TEACHER["handoverScoring"]["headers"], [22, 14, 18, 60, 8, 8, 8, 8, 9, 40], teacher=True)
score_rule = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True, showErrorMessage=True, errorTitle="SBAR score", error="Enter 0, 1 or 2")
scoring.add_data_validation(score_rule)
score_rule.add(f"E2:H{N}")
# Rows (and their Total formula) are appended by the export for each new handover.
for row in sample.get("__handoverScoring", []):
    scoring.append(row)
for letter in "EFGH":
    scoring.column_dimensions[letter].width = 8

for key in ["reviews", "feedback", "roster"]:
    ws = sheet(TABS[key]["name"], TABS[key]["headers"], widths.get(key))
    for row in sample.get(TABS[key]["name"], [])[1:]:
        ws.append(row)

excl = sheet(TEACHER["exclusions"]["name"], TEACHER["exclusions"]["headers"], [16, 50], teacher=True)
excl["D1"] = "Type a Student ID and a reason. The next export marks that student as excluded from all statistics (their rows stay)."
excl["D1"].font = note_font
for row in sample.get("__exclusions", []):
    excl.append(row)

log = sheet(TABS["log"]["name"], TABS["log"]["headers"], widths["log"])
for row in sample.get(TABS["log"]["name"], [])[1:]:
    log.append(row)

wb.save(out_path)
print(f"wrote {out_path}: {', '.join(wb.sheetnames)}")
